#!/usr/bin/env python3
"""把《桥梁病害成因分析与处治建议》Excel 转成 bi-biz 使用的成因/处治建议知识库 JSON。

用法：
    pip install openpyxl
    python scripts/disease_cause_kb_from_excel.py /path/桥梁病害成因分析与处治建议修订V3.xlsx

默认输出 bi-biz/src/main/resources/json/disease_cause_kb.json。
通用耐久性缺损各版本与梁桥 V2 不转换；Excel 改版时只需调整 SHEETS 中的 sheet 名称或列号。
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

DEFAULT_OUTPUT = Path(__file__).resolve().parent.parent / "bi-biz/src/main/resources/json/disease_cause_kb.json"

# 列号从 0 开始；None 表示该 sheet 没有此列。
STANDARD_COLUMNS = dict(structure=1, category=2, component=3, defect_class=4, defect_type=5,
                        defect_name=6, condition=7, cause=8, suggestion=9)
SHEETS = [
    dict(name="上部结构-梁桥 (V3)", code="BEAM", label="上部结构-梁桥", columns=STANDARD_COLUMNS),
    dict(name="上部结构-拱桥", code="ARCH", label="上部结构-拱桥", columns=STANDARD_COLUMNS),
    dict(name="上部结构-组合体系拱桥", code="COMBO_ARCH", label="上部结构-组合体系拱桥", columns=STANDARD_COLUMNS),
    dict(name="上部结构-悬索桥", code="SUSP", label="上部结构-悬索桥", columns=STANDARD_COLUMNS),
    dict(name="上部结构-斜拉桥", code="CABLE", label="上部结构-斜拉桥", columns=STANDARD_COLUMNS),
    dict(name="下部结构", code="SUB", label="下部结构",
         columns=dict(structure=0, category=1, component=2, defect_class=3, defect_type=4,
                      defect_name=5, condition=6, cause=7, suggestion=8)),
    # 附属构造、桥面系使用合并单元格分层，结构/部件/缺损类别/子类型需要向下填充。
    dict(name="附属构造", code="ANC", label="附属构造", fill_down=(0, 1, 2),
         columns=dict(structure=0, category=1, component=1, defect_class=2, defect_type=3,
                      defect_name=None, condition=4, cause=5, suggestion=6)),
    dict(name="桥面系", code="DECK", label="桥面系", fill_down=(0, 1, 2, 3),
         columns=dict(structure=0, category=1, component=3, defect_class=2, defect_type=4,
                      defect_name=None, condition=5, cause=6, suggestion=7)),
]


def find_sheet(workbook, name):
    for sheet in workbook.worksheets:
        if sheet.title.strip() == name:
            return sheet
    raise SystemExit(f"Excel 中找不到 sheet：{name}")


def merged_values(sheet):
    values = {}
    for merged in sheet.merged_cells.ranges:
        value = sheet.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                values[(row, col)] = value
    return values


def name_text(value):
    """名称类字段：去掉换行和所有空白（Excel 中有“锈 蚀”这类写法）。"""
    if value is None:
        return ""
    return re.sub(r"[\s\u3000\u00a0]+", "", str(value))


def long_text(value):
    """成因/处治建议：保留分行，去掉每行首尾空白。"""
    if value is None:
        return ""
    lines = [line.strip() for line in str(value).replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    text = "\n".join(line for line in lines if line)
    # Excel 中残留的 0、z 等单字符不是有效内容。
    return "" if len(text) <= 1 else text


def convert_sheet(sheet, config, warnings):
    columns = config["columns"]
    merged = merged_values(sheet)
    fill_down = config.get("fill_down", ())
    last = {}
    entries = []
    for row_index in range(2, sheet.max_row + 1):
        raw = []
        for col_index in range(sheet.max_column):
            value = sheet.cell(row_index, col_index + 1).value
            if value is None:
                value = merged.get((row_index, col_index + 1))
            raw.append(value)
        for col_index in fill_down:
            if raw[col_index] is None or str(raw[col_index]).strip() == "":
                raw[col_index] = last.get(col_index)
            else:
                last[col_index] = raw[col_index]

        def column(key, text=name_text):
            index = columns[key]
            return "" if index is None or index >= len(raw) else text(raw[index])

        defect_type = column("defect_type")
        defect_name = column("defect_name")
        cause = column("cause", long_text)
        suggestion = column("suggestion", long_text)
        if not defect_type and not defect_name:
            continue
        if not cause and not suggestion:
            warnings.append(f"{config['label']} 第{row_index}行没有成因和处治建议，已跳过")
            continue
        if not cause:
            warnings.append(f"{config['label']} 第{row_index}行成因为空")
        if not suggestion:
            warnings.append(f"{config['label']} 第{row_index}行处治建议为空")
        entries.append({
            "id": f"{config['code']}-{row_index}",
            "sheet": config["label"],
            "structure": column("structure"),
            "category": column("category"),
            "component": column("component"),
            "defectClass": column("defect_class"),
            "defectType": defect_type,
            "defectName": "" if defect_name == defect_type else defect_name,
            "condition": column("condition"),
            "cause": cause,
            "suggestion": suggestion,
        })
    return entries


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.excel, data_only=True)
    warnings = []
    entries = []
    for config in SHEETS:
        sheet_entries = convert_sheet(find_sheet(workbook, config["name"]), config, warnings)
        print(f"{config['label']}: {len(sheet_entries)} 条")
        entries.extend(sheet_entries)

    output = {"source": args.excel.name, "entries": entries}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"共 {len(entries)} 条，已写入 {args.output}")
    for warning in warnings:
        print("警告：" + warning, file=sys.stderr)


if __name__ == "__main__":
    main()
